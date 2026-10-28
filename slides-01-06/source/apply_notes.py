"""apply_notes.py — writes ../SPEAKER_NOTES.md into the PowerPoint's speaker notes.

Usage: python apply_notes.py [deck.pptx]    (default: ../Slides-01-06.pptx)

Each paragraph goes to the PPT slide it names ("Slide 34", "(slide 2)", "slides 13–28" → first number);
paragraphs without a number go to the last slide named in that section. Step slides (2's happy/chaotic
steps) keep their per-step caption + code line underneath.
"""
import os
import re
import sys
from pptx import Presentation

STEP_SLIDES = set(range(3, 12)) | set(range(13, 29))


def parse(md):
    notes = {}
    body = md.split("## If someone asks")[0]
    for sec in re.split(r"\n## Slide ", body)[1:]:
        m = re.search(r"\*PPT slides? (\d+)(?:–(\d+))?", sec)
        cur = int(m.group(1)) if m else None
        last = int(m.group(2)) if m and m.group(2) else cur
        sec = sec.split("\n", 1)[1] if "\n" in sec else ""  # drop the "N — Title" header line
        for para in re.split(r"\n\s*\n", sec):
            p = para.strip()
            if not p or p.startswith("*PPT") or p.startswith("---"):
                continue
            if p.startswith("**"):
                head = p.split(":**")[0]
                n = re.search(r"[Ss]lides? (\d+)", head)
                if n:
                    cur = int(n.group(1))
            if re.match(r"\*\*(Close|Bridge|Hand-over)", p):
                cur = last
            if cur:
                clean = re.sub(r"\*\*|`", "", p)
                notes.setdefault(cur, []).append(clean)
    return {k: "\n\n".join(v) for k, v in notes.items()}


def main(path):
    md = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "SPEAKER_NOTES.md"), encoding="utf-8").read()
    notes = parse(md)
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        tf = slide.notes_slide.notes_text_frame
        old = tf.text.strip()
        new = notes.get(i)
        if new and i in STEP_SLIDES:
            tf.text = new + "\n\n— This step: " + old
        elif new:
            tf.text = new
    prs.save(path)
    print(f"Notes applied to {len(notes)} slides:", sorted(notes))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Slides-01-06.pptx"))
