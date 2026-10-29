"""
make_gif.py — turn a range of PowerPoint slides into one animated GIF.

    python make_gif.py                      # asks for everything
    python make_gif.py --pptx deck.pptx --from 3 --to 8          # or answer up front

It asks for:
    1. the PowerPoint file
    2. the first and last slide to include
    3. where to save the GIF (default: a "gifs" folder next to the PowerPoint)

How it works:
    RENDER  each slide in the range becomes a PNG (PowerPoint on Windows, otherwise LibreOffice)
    STITCH  the PNGs are joined into a looping GIF with a short cross-fade between slides
    SAVE    gifs/<deck>-slides-<from>-<to>.gif

Install:  pip install pillow numpy python-pptx pymupdf
          + Windows with PowerPoint: pip install pywin32      (or LibreOffice on any OS)
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

import numpy as np
from PIL import Image

# ---------------------------------------------------------------- settings (safe to change)
RENDER_WIDTH = 3200     # render big, then shrink = crisp text
GIF_WIDTH = 1600        # final GIF width in pixels
COLORS = 256            # GIF maximum
HOLD_MS = 1300          # time on each slide
FIRST_HOLD_MS = 2000    # first slide stays a little longer
LAST_HOLD_MS = 3000     # last slide stays longest before the loop restarts
FADE_FRAMES = 3         # in-between frames for the cross-fade (0 = hard cut)
FADE_MS = 70            # time per in-between frame


# ---------------------------------------------------------------- small helpers
def ask(question, default=None):
    """Ask in the terminal; Enter keeps the default. Quotes from 'Copy as path' are removed."""
    hint = f" [{default}]" if default not in (None, "") else ""
    answer = input(f"{question}{hint}: ").strip().strip('"').strip("'")
    return answer or (str(default) if default is not None else "")


def slide_count(pptx):
    from pptx import Presentation
    return len(Presentation(pptx).slides)


def slide_width_inches(pptx):
    from pptx import Presentation
    return Presentation(pptx).slide_width / 914400


# ---------------------------------------------------------------- 1 · RENDER
def render(pptx, first, last, work):
    """PNG per slide in [first, last]. Tries PowerPoint (Windows), then LibreOffice."""
    for fn in (render_powerpoint, render_libreoffice):
        try:
            pngs = fn(pptx, first, last, work)
            print(f"  rendered slides {first}–{last} with {fn.__name__.replace('render_', '')}")
            return pngs
        except Exception as e:
            print(f"  ({fn.__name__.replace('render_', '')} not used: {e})")
    sys.exit("Could not render slides. Install PowerPoint + pywin32 (Windows) or LibreOffice + pymupdf.")


def render_powerpoint(pptx, first, last, work):
    if os.name != "nt":
        raise RuntimeError("not Windows")
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(os.path.abspath(pptx), ReadOnly=True, WithWindow=False)
    h = int(RENDER_WIDTH * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
    pngs = []
    for n in range(first, last + 1):
        out = os.path.join(work, f"slide-{n:03d}.png")
        pres.Slides(n).Export(out, "PNG", RENDER_WIDTH, h)
        pngs.append(out)
    pres.Close()
    return pngs


def render_libreoffice(pptx, first, last, work):
    soffice = shutil.which("soffice") or shutil.which("libreoffice") or r"C:\Program Files\LibreOffice\program\soffice.exe"
    if not (shutil.which(soffice) or os.path.exists(soffice)):
        raise RuntimeError("LibreOffice not found")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", work, os.path.abspath(pptx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    pdf = os.path.join(work, os.path.splitext(os.path.basename(pptx))[0] + ".pdf")
    import pymupdf
    doc = pymupdf.open(pdf)
    pngs = []
    for n in range(first, last + 1):
        page = doc[n - 1]
        zoom = RENDER_WIDTH / page.rect.width
        out = os.path.join(work, f"slide-{n:03d}.png")
        page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).save(out)
        pngs.append(out)
    return pngs


# ---------------------------------------------------------------- 2 · STITCH
def load(png, crop_top, crop_bottom, slide_w_in):
    img = Image.open(png).convert("RGB")
    ppi = img.width / slide_w_in
    img = img.crop((0, round(crop_top * ppi), img.width, img.height - round(crop_bottom * ppi)))
    return img.resize((GIF_WIDTH, round(img.height * GIF_WIDTH / img.width)), Image.LANCZOS)


def palette_for(steps):
    """One shared palette from all slides (no flicker); near-white snapped to pure white."""
    w, h = steps[0].size
    sheet = Image.new("RGB", (w, h * len(steps)), "white")
    for i, s in enumerate(steps):
        sheet.paste(s, (0, i * h))
    pal = sheet.quantize(colors=COLORS, method=Image.Quantize.MAXCOVERAGE).getpalette()[: COLORS * 3]
    for c in range(COLORS):
        if min(pal[c * 3: c * 3 + 3]) >= 246:
            pal[c * 3: c * 3 + 3] = [255, 255, 255]
    return pal


def to_indexed(frames, pal):
    """Exact nearest-colour match for every pixel (Pillow's own lookup can drift a few shades)."""
    p = np.array(pal, dtype=np.int32).reshape(-1, 3)
    packed = [(a[..., 0] << 16) | (a[..., 1] << 8) | a[..., 2] for a in (np.asarray(f, dtype=np.int32) for f in frames)]
    uniq = np.unique(np.concatenate([np.unique(x) for x in packed]))
    rgb = np.stack([(uniq >> 16) & 255, (uniq >> 8) & 255, uniq & 255], axis=1)
    lut = np.empty(len(uniq), dtype=np.uint8)
    for s in range(0, len(uniq), 20000):
        lut[s:s + 20000] = ((rgb[s:s + 20000, None, :] - p[None, :, :]) ** 2).sum(axis=2).argmin(axis=1)
    out = []
    for x in packed:
        img = Image.fromarray(lut[np.searchsorted(uniq, x)], mode="P")
        img.putpalette(pal)
        out.append(img)
    return out


def stitch(steps, out_path, hold_ms):
    frames, durations = [], []
    for i, img in enumerate(steps):
        frames.append(img)
        durations.append(FIRST_HOLD_MS if i == 0 else LAST_HOLD_MS if i == len(steps) - 1 else hold_ms)
        nxt = steps[(i + 1) % len(steps)]                 # last fades back to first: smooth loop
        for k in range(1, FADE_FRAMES + 1 if len(steps) > 1 else 1):
            frames.append(Image.blend(img, nxt, k / (FADE_FRAMES + 1)))
            durations.append(FADE_MS)
    gif = to_indexed(frames, palette_for(steps))
    gif[0].save(out_path, save_all=True, append_images=gif[1:], duration=durations, loop=0, optimize=True)
    return sum(durations) / 1000


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Make an animated GIF from a range of PowerPoint slides")
    ap.add_argument("--pptx"); ap.add_argument("--from", dest="first", type=int); ap.add_argument("--to", dest="last", type=int)
    ap.add_argument("--out", help="folder for the GIF"); ap.add_argument("--name", help="GIF file name")
    ap.add_argument("--hold", type=int, default=HOLD_MS, help="ms per slide")
    ap.add_argument("--crop-top", type=float, default=0.0, help="inches to cut from the top (e.g. the title)")
    ap.add_argument("--crop-bottom", type=float, default=0.0, help="inches to cut from the bottom (e.g. the footer)")
    a = ap.parse_args()

    print("\n=== make_gif · PowerPoint slides → animated GIF ===\n")
    pptx = a.pptx or ask("1) PowerPoint file (path)")
    if not os.path.isfile(pptx):
        sys.exit(f"Not found: {pptx}")
    total = slide_count(pptx)
    print(f"   {os.path.basename(pptx)} has {total} slides.")
    first = a.first or int(ask("2) First slide", 1))
    last = a.last or int(ask("   Last slide", total))
    if not (1 <= first <= last <= total):
        sys.exit(f"Slide range must be within 1–{total}, first ≤ last.")
    out_dir = a.out or ask("3) Save the GIF in folder", os.path.join(os.path.dirname(os.path.abspath(pptx)), "gifs"))
    stem = re.sub(r"[^A-Za-z0-9]+", "-", os.path.splitext(os.path.basename(pptx))[0]).strip("-")
    name = a.name or ask("   File name", f"{stem}-slides-{first}-{last}.gif")
    if not name.lower().endswith(".gif"):
        name += ".gif"
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, name)

    print("\nWorking…")
    with tempfile.TemporaryDirectory() as work:
        pngs = render(pptx, first, last, work)
        steps = [load(p, a.crop_top, a.crop_bottom, slide_width_inches(pptx)) for p in pngs]
        secs = stitch(steps, out, a.hold)
    print(f"\n✓ GIF saved: {os.path.abspath(out)}")
    print(f"  {len(steps)} slides · {secs:.0f}s loop · {os.path.getsize(out) / 1e6:.1f} MB\n")


if __name__ == "__main__":
    main()
