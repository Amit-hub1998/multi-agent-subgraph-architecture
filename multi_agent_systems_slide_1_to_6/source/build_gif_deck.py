"""
build_gif_deck.py — build the short deck from the master deck + the GIFs.

    python build_gif_deck.py                                  # → ../Slides-01-06-GIF.pptx
    python build_gif_deck.py --out C:/share/Deck.pptx         # anywhere: GIFs are copied to C:/share/gifs/

HOW IT WORKS
    It opens the 43-slide master deck and edits a copy of it:
      1. ADD     new slides that each show one GIF (+ a "► Play animation" link)
      2. KEEP    the master slides listed in PLAN as they are (their Morph transitions come along)
      3. DROP    every other master slide
      4. ORDER   the slides as in PLAN and rebuild the sections
      5. TWEAK   slide 4 (failure modes): Fade transition + families appear one per click
      6. COPY    the GIFs into a "gifs" folder next to the new deck, so the relative links always work

    Edit the master deck, re-run pptx_to_gifs.py and this script, and the short deck follows.

REQUIREMENTS
    pip install python-pptx
"""

import argparse
import copy
import os
import re
import shutil
import uuid

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))

# --------------------------------------------------------------------------------------------
# The plan: what the short deck contains, in order.
#   ("master", n)            keep slide n of the master deck (1-based)
#   ("gif", {...})           a new slide showing a GIF from the gifs folder
# Sections group the slides in PowerPoint's slide panel.
# --------------------------------------------------------------------------------------------
PLAN = [
    ("1 · Response Codes", [("master", 1)]),
    ("2 · Why Single Agents Fail", [
        ("gif", dict(gif="2-why-single-agents-fail-happy-path.gif", notes="happy",
                     title="Why Single Agents Fail · One Ticket",
                     lead="Agents suffer from overload too. With one BAU ticket, the agent stays focused: right query, right fix, sign-off.",
                     takeaway="Same agent, same five tools: one ticket, fixed and signed off in 4 LLM calls.")),
        ("gif", dict(gif="2-why-single-agents-fail-chaotic-path.gif", notes="chaos",
                     title="Why Single Agents Fail · Three Tickets",
                     lead="Give the same agent three tickets in one session and its decisions get weaker: mixed-up data, repeated retries, a test edited until it passes.",
                     takeaway="A single agent is a monolith: it works for one task and fails like a monolith when you scale the work.")),
    ]),
    ("3 · Why Single Agents Fail – Part 2", [("master", 29), ("master", 30)]),   # mixed → sorted (Morph)
    ("4 · Failure Modes", [("master", 33)]),                                    # one slide, reveal on click
    ("5 · Root Cause", [("master", 34), ("master", 39)]),                       # before → all ticked (Morph)
    ("6 · Core Idea", [
        ("gif", dict(gif="6-core-idea.gif", notes=6,
                     title="The Core Idea of Multi-Agent Systems",
                     lead="Give each agent one job and connect them with clear hand-offs. The shape can change (supervisor, pipeline, swarm); the building blocks stay the same.",
                     takeaway="A multi-agent system is a well-run team: one job each, clear hand-offs, a shared goal, and someone who signs off.")),
    ]),
]

# Slide 4: which shapes appear on each click (matched on the start of the shape name, see the Selection Pane)
REVEAL = {33: [["!!fi1", "!!fh1", "!!card1", "!!ci1", "!!cn1", "!!cb1"],                 # click 1: acting
               ["!!fi2", "!!fh2", "!!card2", "!!ci2", "!!cn2", "!!cb2", "!!take4"]]}     # click 2: checking
FADE_SLIDES = {33}   # use a Fade transition here instead of Morph

# Replace the text of named shapes on kept master slides (formatting is kept)
TEXT_OVERRIDE = {
    34: {"!!pstep": "Before"},
    39: {"!!pstep": "After · all five principles applied"},
}

# Notes for kept master slides that need a fuller note now that the step slides are gone
NOTES_OVERRIDE = {
    33: "section4",
    39: ("All five principles applied: separation of concerns, isolated state, bounded execution, maker-checker, "
         "and explicit hand-offs. Look what we've built: a multi-agent system.\n\n"
         "None of this is new. We don't need a smarter model or a longer prompt. This isn't an AI problem; it's a "
         "system design problem we've already solved, just not for agents yet."),
}

# Design system (LOCKED.md)
TEXT, PURPLE_F, PURPLE_O, PURPLE_T, FRAME = "2C2C2A", "EEEDFE", "534AB7", "3C3489", "D3D1C7"
FONT = "Calibri"          # titles and diagram text
BODY_FONT = "Arial"       # lead line and footer (house style)

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
      "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


# --------------------------------------------------------------------------------------------
# Speaker notes from ../SPEAKER_NOTES.md
# --------------------------------------------------------------------------------------------
def clean_note(text, keep_clicks):
    """Remove the master deck's slide numbers ("Slide 29 (mixed):", "(slides 13–28 …)") from a note."""
    out = []
    for p in re.split(r"\n\s*\n", text):
        p = re.sub(r"\s*\((?:PPT )?slides? [\d–]+[^)]*\)", "", p)                       # "(slides 13–28, click fast)"
        click = p.startswith("[click]")
        p = re.sub(r"^\[click\]\s*", "", p)
        p = re.sub(r"^Slides? \d+(?:–\d+)?\s*(?:—\s*)?", "", p)                        # "Slide 40 — "
        p = re.sub(r"^\((click[^)]*)\):\s*", "", p)                                    # "(click fast …):" no longer applies
        p = re.sub(r"^\(([^)]*)\):", lambda m: m.group(1).capitalize() + ":", p)       # "(mixed):" → "Mixed:"
        p = re.sub(r"^:\s*", "", p)
        out.append(("[click] " if click and keep_clicks else "") + p)
    return "\n\n".join(out)


def load_notes():
    md = open(os.path.join(HERE, "..", "SPEAKER_NOTES.md"), encoding="utf-8").read().split("## If someone asks")[0]
    sections = {}
    for sec in md.split("\n## Slide ")[1:]:
        n = int(re.match(r"\d+", sec).group())
        paras = [p.strip() for p in re.split(r"\n\s*\n", sec.split("\n", 1)[1]) if p.strip() and not p.startswith(("*PPT", "---"))]
        sections[n] = [re.sub(r"\*\*|`", "", p) for p in paras]
    return {
        "happy": "\n\n".join(p for p in sections[2] if p.startswith(("Opening", "Happy flow"))),
        "chaos": "\n\n".join(p for p in sections[2] if not p.startswith(("Opening", "Happy flow"))),
        "section4": "\n\n".join(sections[4]),
        6: "\n\n".join(sections[6]),
    }


# --------------------------------------------------------------------------------------------
# 1 · ADD a GIF slide
# --------------------------------------------------------------------------------------------
def gif_size(path):
    """A GIF file starts with 'GIF89a', then width and height as 2-byte little-endian numbers."""
    with open(path, "rb") as f:
        head = f.read(10)
    return int.from_bytes(head[6:8], "little"), int.from_bytes(head[8:10], "little")


def textbox(slide, text, x, y, w, h, size, bold=False, italic=False, color=TEXT, anchor=MSO_ANCHOR.TOP, font=FONT, align=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold, run.font.italic = font, Pt(size), bold, italic
    run.font.color.rgb = RGBColor.from_string(color)
    if align:
        tf.paragraphs[0].alignment = align
    return tb


def house_style(slide, title, lead):
    """Same header and footer as the master deck: bold title, thin rule, lead line, page number (Office adds the classification label)."""
    textbox(slide, title, 0.45, 0.22, 12.4, 0.58, 26, bold=True, color="1A1A1A", anchor=MSO_ANCHOR.BOTTOM)
    rule = slide.shapes.add_connector(1, Inches(0.45), Inches(0.88), Inches(12.88), Inches(0.88))   # 1 = straight line
    rule.line.color.rgb = RGBColor.from_string("7F7F7F"); rule.line.width = Pt(0.75)
    if lead:
        textbox(slide, lead, 0.45, 0.97, 12.4, 0.48, 14, color="404040", font=BODY_FONT)
    # Page number: a real slide-number field, so it stays right when slides move
    num = textbox(slide, "", 6.17, 6.98, 1.0, 0.22, 10, color="595959", font=BODY_FONT, align=PP_ALIGN.CENTER)
    p = num.text_frame.paragraphs[0]._p
    for r in p.findall(f"{{{NS['a']}}}r"):
        p.remove(r)
    fld = etree.SubElement(p, f"{{{NS['a']}}}fld", id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    rpr = etree.SubElement(fld, f"{{{NS['a']}}}rPr", lang="en-US", sz="1000")
    fill = etree.SubElement(rpr, f"{{{NS['a']}}}solidFill")
    etree.SubElement(fill, f"{{{NS['a']}}}srgbClr", val="595959")
    etree.SubElement(rpr, f"{{{NS['a']}}}latin", typeface=BODY_FONT)
    etree.SubElement(fld, f"{{{NS['a']}}}t").text = "‹#›"


def add_gif_slide(prs, gif_path, link, title, lead, takeaway, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[0])   # the master deck's only layout (no placeholders)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    house_style(slide, title, lead)

    # Fit the GIF into 12.4" × 4.75" below the lead line, keep its shape, centre it
    w0, h0 = gif_size(gif_path)
    k = min(12.4 / w0, 4.75 / h0)
    w, h = w0 * k, h0 * k
    x, y = (prs.slide_width / 914400 - w) / 2, 1.58

    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 0.06), Inches(y - 0.06), Inches(w + 0.12), Inches(h + 0.12))
    frame.adjustments[0] = 0.01
    frame.fill.solid(); frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = RGBColor.from_string(FRAME); frame.line.width = Pt(0.75)
    frame.shadow.inherit = False
    slide.shapes.add_picture(gif_path, Inches(x), Inches(y), Inches(w), Inches(h))   # stored as-is: stays animated

    # "► Play animation" button: the text run carries the hyperlink
    by = y + h + 0.16
    btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(by), Inches(1.9), Inches(0.34))
    btn.adjustments[0] = 0.5
    btn.fill.solid(); btn.fill.fore_color.rgb = RGBColor.from_string(PURPLE_F)
    btn.line.color.rgb = RGBColor.from_string(PURPLE_O); btn.line.width = Pt(1)
    btn.shadow.inherit = False
    tf = btn.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "►  Play animation"
    run.font.name, run.font.size, run.font.bold = FONT, Pt(11), True
    run.font.color.rgb = RGBColor.from_string(PURPLE_T)
    run.hyperlink.address = link            # relative, e.g. "gifs/6-core-idea.gif"

    textbox(slide, takeaway, x + 2.1, by - 0.02, w - 2.1, 0.4, 14, bold=True, italic=True, anchor=MSO_ANCHOR.MIDDLE)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


# --------------------------------------------------------------------------------------------
# 5 · TWEAK slide 4: Fade transition + reveal on click
# --------------------------------------------------------------------------------------------
def set_fade_transition(slide):
    """Replace the slide's transition (Morph) with a plain Fade."""
    sld = slide._element
    for el in sld.findall("mc:AlternateContent", NS) + sld.findall("p:transition", NS):
        sld.remove(el)
    trans = etree.SubElement(sld, f"{{{NS['p']}}}transition", spd="med")
    etree.SubElement(trans, f"{{{NS['p']}}}fade")
    clr = sld.find("p:clrMapOvr", NS)
    clr.addnext(trans)
    return trans


def add_click_reveal(slide, groups, after):
    """Each group of shapes fades in on one click (PowerPoint's 'Fade' entrance effect)."""
    ids = {sh.name: sh for sh in slide.shapes}
    ctr = iter(range(3, 10000))
    P = "{%s}" % NS["p"]

    def el(tag, parent=None, **attrs):
        e = etree.SubElement(parent, P + tag) if parent is not None else etree.Element(P + tag)
        for k, v in attrs.items():
            e.set(k, str(v))
        return e

    timing = el("timing")
    root = el("cTn", el("par", el("tnLst", timing)), id=1, dur="indefinite", restart="never", nodeType="tmRoot")
    seq = el("seq", el("childTnLst", root), concurrent=1, nextAc="seek")
    main = el("cTn", seq, id=2, dur="indefinite", nodeType="mainSeq")
    clicks = el("childTnLst", main)
    built = []
    for group in groups:
        shapes = [sh for name, sh in ids.items() if any(name.startswith(prefix) for prefix in group)]
        click = el("cTn", el("par", clicks), id=next(ctr), fill="hold")
        el("cond", el("stCondLst", click), delay="indefinite")
        inner = el("cTn", el("par", el("childTnLst", click)), id=next(ctr), fill="hold")
        el("cond", el("stCondLst", inner), delay=0)
        effects = el("childTnLst", inner)
        for i, sh in enumerate(shapes):
            eff = el("cTn", el("par", effects), id=next(ctr), presetID=10, presetClass="entr", presetSubtype=0,
                     fill="hold", grpId=0, nodeType="clickEffect" if i == 0 else "withEffect")
            el("cond", el("stCondLst", eff), delay=0)
            beh = el("childTnLst", eff)
            s = el("cBhvr", el("set", beh))
            el("cond", el("stCondLst", el("cTn", s, id=next(ctr), dur=1, fill="hold")), delay=0)
            el("spTgt", el("tgtEl", s), spid=sh.shape_id)
            el("attrName", el("attrNameLst", s)).text = "style.visibility"
            el("strVal", el("to", s.getparent()), val="visible")
            fade = el("cBhvr", el("animEffect", beh, transition="in", filter="fade"))
            el("cTn", fade, id=next(ctr), dur=500)
            el("spTgt", el("tgtEl", fade), spid=sh.shape_id)
            if sh.shape_type != 13:   # pictures don't get a build entry
                built.append(sh.shape_id)
    for tag, evt in (("prevCondLst", "onPrev"), ("nextCondLst", "onNext")):
        el("sldTgt", el("tgtEl", el("cond", el(tag, seq), evt=evt, delay=0)))
    bld = el("bldLst", timing)
    for sid in built:
        el("bldP", bld, spid=sid, grpId=0, animBg=1)
    after.addnext(timing)


# --------------------------------------------------------------------------------------------
# 2–4 · KEEP / DROP / ORDER slides and rebuild sections
# --------------------------------------------------------------------------------------------
def reorder_and_drop(prs, ordered_slides):
    sld_list = prs.slides._sldIdLst
    entries = list(sld_list)
    keep_ids = [s.slide_id for s in ordered_slides]
    lookup = {int(e.get("id")): e for e in entries}
    for e in entries:                                  # DROP: remove from the list and drop the relationship
        if int(e.get("id")) not in keep_ids:
            prs.part.drop_rel(e.get(f"{{{NS['r']}}}id"))
        sld_list.remove(e)
    for sid in keep_ids:                               # ORDER: put the kept ones back in plan order
        sld_list.append(lookup[sid])


def rebuild_sections(prs, plan_ids):
    pres = prs.part._element
    sec_list = pres.find(".//p14:sectionLst", NS)
    for child in list(sec_list):
        sec_list.remove(child)
    for name, ids in plan_ids:
        sec = etree.SubElement(sec_list, f"{{{NS['p14']}}}section", name=name, id="{%s}" % str(uuid.uuid4()).upper())
        lst = etree.SubElement(sec, f"{{{NS['p14']}}}sldIdLst")
        for sid in ids:
            etree.SubElement(lst, f"{{{NS['p14']}}}sldId", id=str(sid))


# --------------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Short deck: master slides + GIF slides")
    ap.add_argument("--master", default=os.path.join(HERE, "..", "Slides-01-06.pptx"))
    ap.add_argument("--gifs", default=os.path.join(HERE, "..", "gifs"), help="where pptx_to_gifs.py wrote the GIFs")
    ap.add_argument("--out", default=os.path.join(HERE, "..", "Slides-01-06-GIF.pptx"))
    args = ap.parse_args()

    out_dir = os.path.dirname(os.path.abspath(args.out))
    link_dir = os.path.join(out_dir, "gifs")                    # links are always "gifs/<file>" next to the deck
    os.makedirs(link_dir, exist_ok=True)
    notes = load_notes()

    prs = Presentation(args.master)
    master = list(prs.slides)                                   # master[0] is slide 1
    ordered, plan_ids, reveal_ids = [], [], set()
    for section, items in PLAN:
        ids = []
        for kind, ref in items:
            if kind == "master":
                slide = master[ref - 1]
                if ref in REVEAL:
                    reveal_ids.add(slide.slide_id)
                if ref in FADE_SLIDES:
                    trans = set_fade_transition(slide)
                    if ref in REVEAL:
                        add_click_reveal(slide, REVEAL[ref], trans)
                for shape in slide.shapes:
                    new_text = TEXT_OVERRIDE.get(ref, {}).get(shape.name)
                    if new_text and shape.has_text_frame:
                        runs = shape.text_frame.paragraphs[0].runs
                        runs[0].text = new_text
                        for extra in runs[1:]:
                            extra.text = ""
                if ref in NOTES_OVERRIDE:
                    val = NOTES_OVERRIDE[ref]
                    slide.notes_slide.notes_text_frame.text = notes.get(val, val)
            else:
                src = os.path.join(args.gifs, ref["gif"])
                if not os.path.exists(src):
                    raise SystemExit(f"Missing {src} — run: python pptx_to_gifs.py")
                dst = os.path.join(link_dir, ref["gif"])
                if os.path.abspath(src) != os.path.abspath(dst):
                    shutil.copy2(src, dst)                      # COPY: GIF sits in gifs/ next to the deck
                slide = add_gif_slide(prs, dst, "gifs/" + ref["gif"], ref["title"], ref.get("lead"), ref["takeaway"], notes[ref["notes"]])
            ordered.append(slide)
            ids.append(slide.slide_id)
        plan_ids.append((section, ids))

    for slide in ordered:                                       # tidy every note (clicks only matter on slide 4)
        tf = slide.notes_slide.notes_text_frame
        tf.text = clean_note(tf.text, keep_clicks=bool(slide.slide_id in reveal_ids))

    reorder_and_drop(prs, ordered)
    rebuild_sections(prs, plan_ids)
    prs.save(args.out)
    print(f"Wrote {args.out}  ·  {len(ordered)} slides  ·  GIF links → {os.path.relpath(link_dir, out_dir)}/")


if __name__ == "__main__":
    main()
