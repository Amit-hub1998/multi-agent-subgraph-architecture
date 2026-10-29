"""add_morph.py — finishes the "Reading the Graph" slides (16b):
  1. button links keep their text colour (no blue, no underline) and the whole button is clickable
  2. adds a Morph transition so the highlight animates

Usage: python add_morph.py [deck.pptx] [first_slide]      (default: ../Slides-14-16.pptx, from slide 5)

PowerPoint (Microsoft 365 / 2019+) plays Morph; older versions fall back to Fade.
Shapes are paired across slides by name ("!!name"), set in build_slides.js.
"""
import os
import re
import shutil
import sys
import zipfile

MORPH = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
    '<p:transition spd="slow" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" p14:dur="700">'
    '<p159:morph option="byObject"/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>'
    '</mc:AlternateContent>'
)


HLINK_CLR = ('<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ext uri="{A12FA001-AC4F-418D-AE19-62706E023703}">'
             '<ahyp:hlinkClr xmlns:ahyp="http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor" val="tx"/></a:ext></a:extLst>')


def tidy_links(path):
    """Slide-jump links: keep the text colour (no blue), no underline, and make the whole button clickable.
    The text link stays too, because PDF exports only keep links on text."""
    from lxml import etree
    from pptx import Presentation
    from pptx.oxml.ns import qn
    prs = Presentation(path)
    slides = list(prs.slides)
    n = 0
    for slide in slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for r in sh.text_frame._txBody.iter(qn("a:r")):
                rpr = r.find(qn("a:rPr"))
                h = rpr.find(qn("a:hlinkClick")) if rpr is not None else None
                if h is None or "hlinksldjump" not in (h.get("action") or ""):
                    continue
                rpr.set("u", "none")
                if h.find(qn("a:extLst")) is None:
                    h.append(etree.fromstring(HLINK_CLR))
                target = slide.part.related_part(h.get(qn("r:id")))
                sh.click_action.target_slide = next(s for s in slides if s.part is target)
                n += 1
    prs.save(path)
    print("Slide links tidied on", n, "buttons")


def main(path, first):
    tidy_links(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", item.filename)
            if m and int(m.group(1)) >= first:
                xml = data.decode("utf-8")
                if "p159:morph" not in xml:
                    xml = xml.replace("</p:clrMapOvr>", "</p:clrMapOvr>" + MORPH, 1)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)
    print("Morph added from slide", first, ":", path)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "..", "Slides-14-16.pptx"),
         int(sys.argv[2]) if len(sys.argv) > 2 else 5)
